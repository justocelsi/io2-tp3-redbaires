"""Arma el Word del TP3A sobre la plantilla de la cátedra · 26/9/2026.

informe-tp3a.md → pandoc con 25C2-GXX-MC-TITULO.docx de referencia → retoques con python-docx.
Correr desde parte-a/:  python armar_docx.py
"""
from pathlib import Path

import docx
import pypandoc
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

AQUI = Path(__file__).parent
PLANTILLA = AQUI / "25C2-GXX-MC-TITULO.docx"   # plantilla de la cátedra: se baja del campus, no se sube
SALIDA = AQUI / "26C2-GXX-MC-RedBaires.docx"

pypandoc.convert_file(str(AQUI / "informe-tp3a.md"), "docx", outputfile=str(SALIDA),
                      extra_args=[f"--reference-doc={PLANTILLA}", f"--resource-path={AQUI}"])

d = docx.Document(SALIDA)

# encabezado de la plantilla: grupo y cuatrimestre
for s in d.sections:
    for p in s.header.paragraphs:
        for r in p.runs:  # la plantilla parte «Grupo xx» y «2025 c2» en varios runs
            r.text = {"xx": "XX", "5": "6", " c": " C"}.get(r.text, r.text)


def bordes(tabla):
    tblPr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "999999")
        b.append(e)
    tblPr.append(b)


for t in d.tables:
    bordes(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(t.rows):
        for celda in fila.cells:
            for p in celda.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0:
                        r.font.bold = True

# texto corrido un punto más chico que la plantilla y sin espacio de sobra
for p in d.paragraphs:
    if p.style.name in ("Body Text", "First Paragraph", "Normal", "Compact"):
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.size = Pt(11)
    if p.style.name in ("Image Caption", "Caption"):
        for r in p.runs:
            r.font.size = Pt(9)

d.save(SALIDA)
print("escrito", SALIDA)
